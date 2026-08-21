from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd


# ============================================================
# LOOKUP SOURCE STATUS
# ============================================================

@dataclass
class LookupSourceStatus:

    name: str

    path: Path

    exists: bool

    has_data: bool

    rows: int

    sheets: list[str]

    message: str


# ============================================================
# EXCEL CHECK
# ============================================================

def inspect_excel_file(
    name: str,
    path: Path,
) -> LookupSourceStatus:

    # --------------------------------------------------------
    # File doesn't exist
    # --------------------------------------------------------

    if not path.exists():

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=False,
            has_data=False,
            rows=0,
            sheets=[],
            message="File not found",
        )

    # --------------------------------------------------------
    # Try opening workbook
    # --------------------------------------------------------

    try:

        excel = pd.ExcelFile(
            path
        )

    except Exception as exc:

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=True,
            has_data=False,
            rows=0,
            sheets=[],
            message=(
                "Unable to read workbook: "
                f"{exc}"
            ),
        )

    sheets = excel.sheet_names

    if not sheets:

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=True,
            has_data=False,
            rows=0,
            sheets=[],
            message="Workbook contains no sheets",
        )

    # --------------------------------------------------------
    # Check every sheet
    # --------------------------------------------------------

    total_rows = 0

    non_empty_sheets = 0

    for sheet in sheets:

        try:

            df = pd.read_excel(
                path,
                sheet_name=sheet,
            )

        except Exception:

            continue

        if not df.empty:

            # Count rows containing at least
            # one non-null value.

            useful_rows = (
                df
                .dropna(
                    how="all"
                )
                .shape[0]
            )

            if useful_rows > 0:

                non_empty_sheets += 1

                total_rows += useful_rows

    if non_empty_sheets == 0:

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=True,
            has_data=False,
            rows=0,
            sheets=sheets,
            message="Workbook is empty",
        )

    return LookupSourceStatus(
        name=name,
        path=path,
        exists=True,
        has_data=True,
        rows=total_rows,
        sheets=sheets,
        message=(
            f"{total_rows} data rows found"
        ),
    )


# ============================================================
# DOCX CHECK
# ============================================================

def inspect_docx_file(
    name: str,
    path: Path,
) -> LookupSourceStatus:

    if not path.exists():

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=False,
            has_data=False,
            rows=0,
            sheets=[],
            message="File not found",
        )

    try:

        from docx import Document

        document = Document(
            path
        )

    except Exception as exc:

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=True,
            has_data=False,
            rows=0,
            sheets=[],
            message=(
                "Unable to read document: "
                f"{exc}"
            ),
        )

    paragraphs = [
        paragraph.text.strip()
        for paragraph
        in document.paragraphs
        if paragraph.text.strip()
    ]

    table_rows = 0

    for table in document.tables:

        table_rows += len(
            table.rows
        )

    total_content = (
        len(paragraphs)
        + table_rows
    )

    if total_content == 0:

        return LookupSourceStatus(
            name=name,
            path=path,
            exists=True,
            has_data=False,
            rows=0,
            sheets=[],
            message="Document is empty",
        )

    return LookupSourceStatus(
        name=name,
        path=path,
        exists=True,
        has_data=True,
        rows=total_content,
        sheets=[],
        message=(
            f"{total_content} content items found"
        ),
    )


# ============================================================
# INSPECT ALL PART 1 SOURCES
# ============================================================

def inspect_part1_sources(
    paths: dict[str, Path],
) -> dict[str, LookupSourceStatus]:

    result = {}

    for name, path in paths.items():

        if path.suffix.lower() == ".docx":

            result[name] = (
                inspect_docx_file(
                    name,
                    path,
                )
            )

        elif path.suffix.lower() in {
            ".xlsx",
            ".xls",
        }:

            result[name] = (
                inspect_excel_file(
                    name,
                    path,
                )
            )

        else:

            result[name] = LookupSourceStatus(
                name=name,
                path=path,
                exists=path.exists(),
                has_data=path.exists(),
                rows=0,
                sheets=[],
                message="Unsupported source type",
            )

    return result